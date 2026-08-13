import io
import time
from datetime import date

import pytest
from urllib3.exceptions import MaxRetryError

from app.database import Database
from app.main import create_app
from app.models import (
    Application,
    Base,
    Document,
    DocumentJob,
    DocumentOutput,
    ProcessingStep,
    User,
)
from app.parsing import Analysis
from app.worker import process_one


class Objects:
    def open(self, key: str):
        return io.BytesIO(b"%PDF-1.7")


@pytest.fixture
def queued_job() -> tuple[Database, str]:
    app = create_app("sqlite+pysqlite:///:memory:", object_store=Objects())
    Base.metadata.create_all(app.state.database.engine)
    with app.state.database.session() as db:
        user = User(username="owner", password_hash="hash", role="approval_officer")
        db.add(user)
        db.flush()
        application = Application(
            borrower_type="corporate",
            borrower_name="示例企业",
            product="经营贷",
            application_date=date(2026, 8, 7),
            owner_id=user.id,
        )
        db.add(application)
        db.flush()
        document = Document(
            application_id=application.id,
            filename="sample.pdf",
            extension=".pdf",
            declared_mime="application/pdf",
            size_bytes=8,
            sha256="a" * 64,
            object_key="sample",
        )
        db.add(document)
        db.flush()
        job = DocumentJob(document_id=document.id)
        job.steps.append(ProcessingStep(name="validation"))
        db.add(job)
        db.commit()
        return app.state.database, job.id


def test_worker_persists_a_versioned_reviewable_pdf_output() -> None:
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page(width=300, height=200)
    page.insert_text((30, 50), "Reviewable synthetic text")
    content = pdf.tobytes()

    class PdfObjects:
        def open(self, key: str):
            return io.BytesIO(content)

    class Engine:
        version = "test-model-1"

        def analyze(self, image: bytes, *, run_ocr: bool) -> Analysis:
            assert not run_ocr
            return Analysis()

    app = create_app("sqlite+pysqlite:///:memory:", object_store=PdfObjects())
    Base.metadata.create_all(app.state.database.engine)
    with app.state.database.session() as db:
        user = User(username="parser-owner", password_hash="hash", role="approval_officer")
        db.add(user)
        db.flush()
        application = Application(
            borrower_type="corporate",
            borrower_name="解析企业",
            product="经营贷",
            application_date=date(2026, 8, 7),
            owner_id=user.id,
        )
        db.add(application)
        db.flush()
        document = Document(
            application_id=application.id,
            filename="reviewable.pdf",
            extension=".pdf",
            declared_mime="application/pdf",
            size_bytes=len(content),
            sha256="b" * 64,
            object_key="reviewable",
        )
        db.add(document)
        db.flush()
        job = DocumentJob(document_id=document.id)
        for name in ("validation", "parsing_ocr", "seal_detection"):
            job.steps.append(ProcessingStep(name=name, status="waiting"))
        db.add(job)
        db.commit()

    assert process_one(app.state.database, PdfObjects(), "parser-worker", Engine())
    with app.state.database.session() as db:
        output = db.query(DocumentOutput).one()
        assert (output.version, output.status, output.model_version) == (
            1,
            "success",
            "test-model-1",
        )
        assert output.pages[0].blocks[0].text == "Reviewable synthetic text"
        assert {step.name: step.status for step in output.document.jobs[0].steps} == {
            "validation": "success",
            "parsing_ocr": "success",
            "seal_detection": "success",
        }


def test_partial_parse_keeps_validation_success_and_marks_steps_partial(caplog) -> None:
    import pymupdf

    from app.parsing import Analysis

    pdf = pymupdf.open()
    page = pdf.new_page(width=300, height=200)
    page.insert_text((30, 50), "First page stays reviewable")
    pdf.new_page(width=300, height=200)
    content = pdf.tobytes()

    class PdfObjects:
        def open(self, key: str):
            return io.BytesIO(content)

    class PartialEngine:
        version = "test-model-1"

        def analyze(self, image: bytes, *, run_ocr: bool) -> Analysis:
            if run_ocr:
                raise RuntimeError("synthetic OCR failure")
            return Analysis()

    app = create_app("sqlite+pysqlite:///:memory:", object_store=PdfObjects())
    Base.metadata.create_all(app.state.database.engine)
    with app.state.database.session() as db:
        user = User(username="partial-owner", password_hash="hash", role="approval_officer")
        db.add(user)
        db.flush()
        application = Application(
            borrower_type="corporate",
            borrower_name="部分解析企业",
            product="经营贷",
            application_date=date(2026, 8, 7),
            owner_id=user.id,
        )
        db.add(application)
        db.flush()
        document = Document(
            application_id=application.id,
            filename="partial.pdf",
            extension=".pdf",
            declared_mime="application/pdf",
            size_bytes=len(content),
            sha256="c" * 64,
            object_key="partial",
        )
        db.add(document)
        db.flush()
        job = DocumentJob(document_id=document.id)
        for name in ("validation", "parsing_ocr", "seal_detection"):
            job.steps.append(ProcessingStep(name=name, status="waiting"))
        db.add(job)
        db.commit()

    assert process_one(app.state.database, PdfObjects(), "partial-worker", PartialEngine())
    with app.state.database.session() as db:
        job = db.get(DocumentJob, job.id)
        output = db.query(DocumentOutput).one()
        assert (output.status, output.version) == ("partial_success", 1)
        assert [page.status for page in output.pages] == ["success", "failed"]
        assert output.pages[0].blocks[0].text == "First page stays reviewable"
        assert output.document.processing_status == "partial_success"
        assert output.document.review_status == "pending_review"
        steps = {step.name: (step.status, step.error_code) for step in job.steps}
        assert steps["validation"] == ("success", None)
        assert steps["parsing_ocr"] == ("partial_success", "partial_page_failure")
        assert steps["seal_detection"] == ("partial_success", "partial_page_failure")
        assert "First page stays reviewable" not in caplog.text


def test_all_pages_failed_parse_marks_steps_failed() -> None:
    import pymupdf

    from app.parsing import Analysis

    pdf = pymupdf.open()
    pdf.new_page(width=300, height=200)
    content = pdf.tobytes()

    class PdfObjects:
        def open(self, key: str):
            return io.BytesIO(content)

    class FailingEngine:
        version = "test-model-1"

        def analyze(self, image: bytes, *, run_ocr: bool) -> Analysis:
            raise RuntimeError("synthetic total failure")

    app = create_app("sqlite+pysqlite:///:memory:", object_store=PdfObjects())
    Base.metadata.create_all(app.state.database.engine)
    with app.state.database.session() as db:
        user = User(username="failed-owner", password_hash="hash", role="approval_officer")
        db.add(user)
        db.flush()
        application = Application(
            borrower_type="corporate",
            borrower_name="全失败企业",
            product="经营贷",
            application_date=date(2026, 8, 7),
            owner_id=user.id,
        )
        db.add(application)
        db.flush()
        document = Document(
            application_id=application.id,
            filename="all-failed.pdf",
            extension=".pdf",
            declared_mime="application/pdf",
            size_bytes=len(content),
            sha256="d" * 64,
            object_key="all-failed",
        )
        db.add(document)
        db.flush()
        job = DocumentJob(document_id=document.id)
        for name in ("validation", "parsing_ocr", "seal_detection"):
            job.steps.append(ProcessingStep(name=name, status="waiting"))
        db.add(job)
        db.commit()

    assert process_one(app.state.database, PdfObjects(), "failed-worker", FailingEngine())
    with app.state.database.session() as db:
        job = db.get(DocumentJob, job.id)
        assert job.status == "failed"
        assert job.error_code == "all_pages_failed"
        steps = {step.name: (step.status, step.error_code) for step in job.steps}
        assert steps["validation"] == ("success", None)
        assert steps["parsing_ocr"] == ("failed", "all_pages_failed")
        assert job.document.processing_status == "failed"
        assert job.document.review_status == "not_ready"


def test_validation_renews_lease_while_computing(queued_job, monkeypatch) -> None:
    database, job_id = queued_job

    def slow_validate(document, stream) -> None:
        with database.session() as db:
            job = db.get(DocumentJob, job_id)
            before = job.claimed_at
        time.sleep(0.01)
        with database.session() as db:
            job = db.get(DocumentJob, job_id)
            assert job.claimed_at >= before

    monkeypatch.setattr("app.worker.LEASE_RENEWAL_SECONDS", 0.001)
    monkeypatch.setattr("app.worker.validate", slow_validate)
    assert process_one(database, Objects(), "worker")


def test_minio_network_error_is_transient(queued_job, monkeypatch) -> None:
    database, job_id = queued_job

    class UnavailableObjects:
        def open(self, key: str):
            raise MaxRetryError(None, key, "unavailable")

    assert process_one(database, UnavailableObjects(), "worker")
    with database.session() as db:
        job = db.get(DocumentJob, job_id)
        assert job.status == "waiting"
        assert job.error_code == "object_store_unavailable"


def test_unexpected_validator_error_is_permanent_not_transient(queued_job, monkeypatch) -> None:
    database, job_id = queued_job

    def broken(document, stream) -> None:
        raise RuntimeError("programming error")

    monkeypatch.setattr("app.worker.validate", broken)
    assert process_one(database, Objects(), "worker")
    with database.session() as db:
        job = db.get(DocumentJob, job_id)
        assert job.status == "failed"
        assert job.error_code == "unexpected_validation_error"
        assert job.attempts == 1


def test_worker_parses_structured_xlsx_without_image_engine() -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "流水明细"
    sheet.append(["日期", "金额"])
    sheet.append(["2026-08-01", 1234.5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()

    class XlsxObjects:
        def open(self, key: str):
            return io.BytesIO(content)

    app = create_app("sqlite+pysqlite:///:memory:", object_store=XlsxObjects())
    Base.metadata.create_all(app.state.database.engine)
    with app.state.database.session() as db:
        user = User(username="xlsx-owner", password_hash="hash", role="approval_officer")
        db.add(user)
        db.flush()
        application = Application(
            borrower_type="corporate",
            borrower_name="流水企业",
            product="经营贷",
            application_date=date(2026, 8, 7),
            owner_id=user.id,
        )
        db.add(application)
        db.flush()
        document = Document(
            application_id=application.id,
            filename="statement.xlsx",
            extension=".xlsx",
            declared_mime=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            size_bytes=len(content),
            sha256="e" * 64,
            object_key="statement",
        )
        db.add(document)
        db.flush()
        job = DocumentJob(document_id=document.id)
        for name in ("validation", "parsing_ocr"):
            job.steps.append(ProcessingStep(name=name, status="waiting"))
        job.steps.append(ProcessingStep(name="seal_detection", status="not_applicable"))
        db.add(job)
        db.commit()

    assert process_one(app.state.database, XlsxObjects(), "xlsx-worker")
    with app.state.database.session() as db:
        output = db.query(DocumentOutput).one()
        assert (output.version, output.status) == (1, "success")
        assert output.model_version == "none"
        page = output.pages[0]
        assert page.number is None
        assert page.width is None
        block = page.blocks[0]
        assert block.kind == "table"
        assert block.extraction_method == "xlsx_text"
        assert block.locator["sheet"] == "流水明细"
        assert block.locator["cell_range"] == "A1:B2"
        assert block.cells[0].locator["cell"] == "A1"
        assert block.cells[3].text == "1234.5"
        steps = {step.name: step.status for step in output.document.jobs[0].steps}
        assert steps == {
            "validation": "success",
            "parsing_ocr": "success",
            "seal_detection": "not_applicable",
        }
        assert output.document.processing_status == "success"
        assert output.document.review_status == "pending_review"
